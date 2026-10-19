from openpyxl import load_workbook
from collections import Counter

filename = "PTAC_MoA_Reconciliation_Audit_2026-10.xlsx"
wb = load_workbook(filename, read_only=True, data_only=True)

ws = wb["Review Flags"]
counts = Counter()
for row in ws.iter_rows(min_row=2, values_only=True):
    flag = row[7] if len(row) > 7 else None
    if flag:
        counts[str(flag)] += 1

print("\n=== MOST FREQUENT REVIEW FLAGS ===")
for flag, count in counts.most_common(15):
    print(f"{count:5} | {flag}")

for sheet_name, limit in [
    ("Duplicate Registrations", 15),
    ("Mapping Only", 10),
    ("Update Log Review", 12),
]:
    ws = wb[sheet_name]
    print(f"\n=== {sheet_name}: SAMPLE ROWS ===")
    for row in ws.iter_rows(min_row=2, max_row=limit + 1, values_only=True):
        print(" | ".join("" if v is None else str(v) for v in row))

print("\n=== SOURCE ROW COUNTS ===")
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"{sheet_name}: {max(0, ws.max_row - 1)} data rows")

wb.close()
print("\nRead-only review complete. Audit workbook unchanged.")
