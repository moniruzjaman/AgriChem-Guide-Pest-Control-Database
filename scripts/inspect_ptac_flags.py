from openpyxl import load_workbook
from collections import Counter

f = "PTAC_MoA_Reconciliation_Audit_v2.xlsx"
wb = load_workbook(f, read_only=True, data_only=True)
ws = wb["Review Flags"]

rows = list(ws.iter_rows(min_row=2, values_only=True))

def show(title, selected, limit):
    print("\n===", title, "===")
    for r in selected[:limit]:
        print(" | ".join("" if v is None else str(v) for v in r))

show("MOA CODE MISMATCHES",
     [r for r in rows if len(r) > 7 and r[7] == "MOA CODE MISMATCH"], 20)

show("MASTER MOA MISSING/PENDING",
     [r for r in rows if len(r) > 7 and r[7] == "MASTER MOA MISSING OR PENDING"], 15)

show("MAPPING MOA MISSING/PENDING",
     [r for r in rows if len(r) > 7 and r[7] == "MAPPING MOA MISSING OR PENDING"], 15)

for name in ["Unpaired Mapping Rows", "Duplicate Registrations"]:
    sh = wb[name]
    print("\n===", name, "===")
    limit = 100 if name == "Unpaired Mapping Rows" else 20
    for r in sh.iter_rows(min_row=2, max_row=limit+1, values_only=True):
        print(" | ".join("" if v is None else str(v) for v in r))

wb.close()
print("\nRead-only inspection complete; no files changed.")
