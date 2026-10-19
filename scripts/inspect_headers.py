from openpyxl import load_workbook

files = {
    "PTAC_Pesticide_MoA_Master_Database_2026_Rev4.xlsx": [
        "Product_Register", "AI_Dictionary", "MoA_Update_Log"
    ],
    "MoA_Mapping_PTAC_Bangladesh_PolicyEnhanced_CSV-Enriched.xlsx": [
        "Product MoA Mapping", "CSV Safety Data"
    ]
}

for filename, sheets in files.items():
    print("\n" + "=" * 60)
    print(filename)
    wb = load_workbook(filename, read_only=True, data_only=True)

    for name in sheets:
        if name not in wb.sheetnames:
            continue

        ws = wb[name]
        print("\nSHEET:", name)

        for n in range(1, min(ws.max_row, 7) + 1):
            row = next(ws.iter_rows(
                min_row=n, max_row=n, values_only=True
            ))
            values = [
                f"{i+1}={str(v).strip()}"
                for i, v in enumerate(row)
                if v is not None and str(v).strip()
            ]
            if values:
                print("ROW", n, " | ".join(values))

    wb.close()

print("\nDone. Original workbooks were not modified.")
