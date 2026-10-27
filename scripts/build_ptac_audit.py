from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from collections import defaultdict, Counter
from datetime import datetime
import os, re

MASTER_FILE = "PTAC_Pesticide_MoA_Master_Database_2026_Rev4.xlsx"
MAPPING_FILE = "MoA_Mapping_PTAC_Bangladesh_PolicyEnhanced_CSV-Enriched.xlsx"
OUTPUT_FILE = "PTAC_MoA_Reconciliation_Audit_2026-10.xlsx"

for f in (MASTER_FILE, MAPPING_FILE):
    if not os.path.isfile(f):
        raise SystemExit("Required source file not found: " + f)

def norm(value):
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]", "", str(value).lower().strip())

def display(value):
    if value is None:
        return ""
    return str(value).strip()

def read_records(filename, sheet_name, header_row, first_data_row):
    wb = load_workbook(filename, read_only=True, data_only=True)
    ws = wb[sheet_name]
    raw_headers = next(ws.iter_rows(
        min_row=header_row, max_row=header_row, values_only=True
    ))
    headers = {}
    for idx, value in enumerate(raw_headers):
        if value is not None and str(value).strip():
            headers[str(value).strip()] = idx

    records = []
    for row_num, row in enumerate(
        ws.iter_rows(min_row=first_data_row, values_only=True),
        start=first_data_row
    ):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        record = {h: (row[i] if i < len(row) else None)
                  for h, i in headers.items()}
        record["_source_row"] = row_num
        records.append(record)
    wb.close()
    return records, headers

def get(record, *names):
    for name in names:
        if name in record:
            return record[name]
    return None

def reg_key(record, *fields):
    return norm(get(record, *fields))

master, mh = read_records(
    MASTER_FILE, "Product_Register", 5, 6
)
mapping, ph = read_records(
    MAPPING_FILE, "Product MoA Mapping", 6, 7
)
updates, uh = read_records(
    MASTER_FILE, "MoA_Update_Log", 5, 6
)
dictionary, dh = read_records(
    MASTER_FILE, "AI_Dictionary", 5, 6
)
csv_rows, ch = read_records(
    MAPPING_FILE, "CSV Safety Data", 6, 7
)

def index_by_registration(records, *reg_fields):
    result = defaultdict(list)
    for r in records:
        key = reg_key(r, *reg_fields)
        if key:
            result[key].append(r)
    return result

mi = index_by_registration(master, "Registration_No")
pi = index_by_registration(mapping, "Registration No", "Registration_No")
ui = index_by_registration(updates, "Reg_No", "Registration No")

def brand(r, master_side):
    return display(get(r, "Brand_Name" if master_side else "Brand Name",
                      "Brand_Name", "Brand Name"))

def moa(r, master_side):
    return display(get(r, "MoA_Code" if master_side else "MoA Code",
                      "MoA_Code", "MoA Code"))

def common(r, master_side):
    return display(get(r, "Common_Name" if master_side else "Common Name",
                      "Common_Name", "Common Name"))

def is_pending(value):
    v = norm(value)
    return not v or v in ("", "unknown", "pending", "na", "none") or v in (
        "mdash", "notassigned"
    ) or display(value) in ("—", "-", "N/A")

# Summary counts
counts = Counter()
counts["master_records"] = len(master)
counts["mapping_records"] = len(mapping)
counts["ai_dictionary_entries"] = len(dictionary)
counts["update_log_entries"] = len(updates)
counts["csv_safety_rows"] = len(csv_rows)
counts["master_duplicate_registration_groups"] = sum(
    1 for k, v in mi.items() if len(v) > 1
)
counts["mapping_duplicate_registration_groups"] = sum(
    1 for k, v in pi.items() if len(v) > 1
)

# Reconcile master against mapping using registration number first.
reconciliation = []
issues = []
seen_master_regs = Counter()
seen_mapping_regs = Counter()

for r in master:
    reg = display(get(r, "Registration_No"))
    rk = norm(reg)
    if rk:
        seen_master_regs[rk] += 1

    candidates = pi.get(rk, []) if rk else []
    match = candidates[0] if len(candidates) == 1 else None
    flags = []

    if not rk:
        flags.append("MASTER REGISTRATION NUMBER MISSING")
    elif len(candidates) == 0:
        flags.append("NOT FOUND IN MAPPING WORKBOOK")
    elif len(candidates) > 1:
        flags.append("AMBIGUOUS: DUPLICATE REGISTRATION IN MAPPING")

    if rk and len(mi.get(rk, [])) > 1:
        flags.append("DUPLICATE REGISTRATION IN MASTER")

    mbrand = brand(r, True)
    mcommon = common(r, True)
    mmoa = moa(r, True)

    pbrand = brand(match, False) if match else ""
    pcommon = common(match, False) if match else ""
    pmoa = moa(match, False) if match else ""

    if match:
        if norm(mbrand) != norm(pbrand):
            flags.append("BRAND NAME MISMATCH")
        if norm(mcommon) != norm(pcommon):
            flags.append("COMMON NAME MISMATCH")
        if norm(mmoa) != norm(pmoa):
            flags.append("MOA CODE MISMATCH")
        if norm(display(get(r, "Registration_Holder"))) != norm(
            display(get(match, "Registration Holder"))
        ):
            flags.append("REGISTRATION HOLDER MISMATCH")

    if is_pending(mmoa):
        flags.append("MASTER MOA MISSING OR PENDING")
    if match and is_pending(pmoa):
        flags.append("MAPPING MOA MISSING OR PENDING")

    if flags:
        counts["master_rows_with_review_flags"] += 1
        for flag in flags:
            counts[flag] += 1
            issues.append([
                reg, mbrand, mcommon, mmoa, pmoa,
                "Product_Register row " + str(r["_source_row"]),
                "Mapping row " + str(match["_source_row"]) if match else "",
                flag
            ])

    reconciliation.append([
        reg, mbrand, mcommon, mmoa,
        display(get(r, "Scheme")), display(get(r, "Group_No")),
        pbrand, pcommon, pmoa,
        display(get(match, "Scheme")) if match else "",
        display(get(match, "Group No.")) if match else "",
        "MATCHED BY REGISTRATION" if match else "UNMATCHED / AMBIGUOUS",
        "; ".join(flags),
        r["_source_row"],
        match["_source_row"] if match else ""
    ])

# Find mapping-only products and duplicate registrations.
mapping_only = []
for r in mapping:
    reg = display(get(r, "Registration No", "Registration_No"))
    rk = norm(reg)
    if rk:
        seen_mapping_regs[rk] += 1
    if not rk or rk not in mi:
        mapping_only.append([
            reg, brand(r, False), common(r, False), moa(r, False),
            r["_source_row"], "MAPPING RECORD NOT FOUND IN MASTER"
        ])

duplicates = []
for side, index, brand_field in [
    ("MASTER", mi, "Brand_Name"),
    ("MAPPING", pi, "Brand Name")
]:
    for rk, records in index.items():
        if len(records) > 1:
            for r in records:
                duplicates.append([
                    side, display(get(r, "Registration_No", "Registration No")),
                    brand(r, side == "MASTER"), common(r, side == "MASTER"),
                    moa(r, side == "MASTER"), r["_source_row"],
                    "Duplicate normalized registration number; manual review required"
                ])

# Compare update log's new MoA against master records.
update_review = []
for u in updates:
    reg = display(get(u, "Reg_No"))
    rk = norm(reg)
    new_moa = display(get(u, "New_MoA"))
    candidates = mi.get(rk, []) if rk else []
    if not candidates:
        status = "REGISTRATION NOT FOUND IN MASTER"
        current = ""
        bname = display(get(u, "Brand_Name"))
    else:
        candidate = candidates[0]
        current = moa(candidate, True)
        bname = brand(candidate, True)
        if len(candidates) > 1:
            status = "DUPLICATE MASTER REGISTRATION - REVIEW"
        elif norm(current) == norm(new_moa):
            status = "LOG NEW MOA MATCHES MASTER"
        else:
            status = "LOG NEW MOA DIFFERS FROM MASTER - REVIEW"
    update_review.append([
        reg, display(get(u, "Brand_Name")), display(get(u, "Common_Name")),
        display(get(u, "Old_MoA")), new_moa, current, status,
        u["_source_row"]
    ])

# Check curated CSV rows against the register using registration number.
csv_review = []
for r in csv_rows:
    reg = display(get(r, "Registration No"))
    rk = norm(reg)
    candidates = mi.get(rk, []) if rk else []
    csv_brand = display(get(r, "Trade Name"))
    if not rk:
        status = "MISSING REGISTRATION NUMBER"
    elif not candidates:
        status = "NOT FOUND BY REGISTRATION NUMBER"
    elif len(candidates) > 1:
        status = "AMBIGUOUS DUPLICATE REGISTRATION"
    elif norm(csv_brand) != norm(brand(candidates[0], True)):
        status = "REGISTRATION FOUND; BRAND NAME DIFFERS"
    else:
        status = "REGISTRATION AND BRAND MATCH"
    csv_review.append([
        display(get(r, "Source ID")), reg, csv_brand,
        display(get(r, "Common Name")),
        display(get(r, "WHO Toxicity Class")),
        display(get(r, "PHI (days)")), display(get(r, "REI (hrs)")),
        display(get(r, "In Register?")), status, r["_source_row"]
    ])

# Output workbook
out = Workbook()
summary = out.active
summary.title = "Audit Summary"
summary.append(["PTAC Pesticide MoA Reconciliation Audit"])
summary.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
summary.append(["Purpose", "Cross-workbook consistency checks; not official validation"])
summary.append(["Source master", MASTER_FILE])
summary.append(["Source mapping", MAPPING_FILE])
summary.append([])
summary.append(["Metric", "Count / Result"])
for key, value in counts.items():
    summary.append([key, value])
summary.append([])
summary.append(["IMPORTANT LIMITATION"])
summary.append(["A match between workbooks does not prove MoA correctness or current registration status."])
summary.append(["All flags require review against authoritative DAE and IRAC/FRAC/HRAC sources."])

def add_sheet(title, headers, rows):
    ws = out.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    return ws

add_sheet(
    "Product Reconciliation",
    ["Registration No", "Master Brand", "Master Common Name", "Master MoA",
     "Master Scheme", "Master Group", "Mapping Brand", "Mapping Common Name",
     "Mapping MoA", "Mapping Scheme", "Mapping Group", "Match Status",
     "Review Flags", "Master Source Row", "Mapping Source Row"],
    reconciliation
)
add_sheet(
    "Review Flags",
    ["Registration No", "Master Brand", "Master Common Name", "Master MoA",
     "Mapping MoA", "Master Source", "Mapping Source", "Review Flag"],
    issues
)
add_sheet(
    "Mapping Only",
    ["Registration No", "Brand", "Common Name", "MoA Code", "Source Row", "Flag"],
    mapping_only
)
add_sheet(
    "Duplicate Registrations",
    ["Workbook", "Registration No", "Brand", "Common Name", "MoA Code",
     "Source Row", "Review Note"],
    duplicates
)
add_sheet(
    "Update Log Review",
    ["Registration No", "Log Brand", "Log Common Name", "Old MoA",
     "Logged New MoA", "Current Master MoA", "Review Status", "Log Source Row"],
    update_review
)
add_sheet(
    "CSV Safety Review",
    ["Source ID", "Registration No", "Trade Name", "Common Name",
     "WHO Toxicity Class", "PHI Days", "REI Hours", "Original Match Flag",
     "Recomputed Registration/Brand Check", "CSV Source Row"],
    csv_review
)
add_sheet(
    "AI Dictionary",
    ["Active Ingredient", "MoA Code", "Mechanism Target",
     "Resistance Risk", "Rotation Partners", "Source Row"],
    [[display(get(r, "Active_Ingredient")),
      display(get(r, "MoA_Code")),
      display(get(r, "Mechanism_Target")),
      display(get(r, "Resistance_Risk")),
      display(get(r, "Rotation_Partners")),
      r["_source_row"]] for r in dictionary]
)

# Style every sheet for readable review.
for ws in out.worksheets:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="006A4E")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 12
        for cell in ws.iter_rows(min_col=col, max_col=col):
            value = cell[0].value
            if value is not None:
                max_len = max(max_len, min(len(str(value)), 42))
        ws.column_dimensions[letter].width = min(max_len + 2, 44)
    ws.sheet_view.showGridLines = False

out.save(OUTPUT_FILE)

print("\nAUDIT WORKBOOK CREATED:", OUTPUT_FILE)
print("Master records:", len(master))
print("Mapping records:", len(mapping))
print("Product rows flagged for review:", counts["master_rows_with_review_flags"])
print("Mapping-only records:", len(mapping_only))
print("Duplicate registration rows:", len(duplicates))
print("Update log rows reviewed:", len(update_review))
print("CSV safety rows reviewed:", len(csv_review))
print("\nOriginal workbooks were not modified.")
print("This is a consistency audit, not official MoA validation.")
