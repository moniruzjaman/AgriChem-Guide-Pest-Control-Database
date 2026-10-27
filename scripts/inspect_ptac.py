from openpyxl import load_workbook
import os

files = [
    (
        "PTAC_Pesticide_MoA_Master_Database_2026_Rev4.xlsx",
        ["Product_Register", "AI_Dictionary", "MoA_Update_Log"]
    ),
    (
        "MoA_Mapping_PTAC_Bangladesh_PolicyEnhanced_CSV-Enriched.xlsx",
        ["Product MoA Mapping", "CSV Safety Data"]
    )
]

for filename, sheets in files:
    print("\n" + "=" * 70)
    print("FILE:", filename)

    if not os.path.exists(filename):
        print("ERROR: File not found.")
        continue

    wb = load_workbook(filename, read_only=True, data_only=True)

    for sheet_name in sheets:
        if sheet_name not in wb.sheetnames:
            print("\nMissing sheet:", sheet_name)
            continue

        ws = wb[sheet_name]
        print("\nSHEET:", sheet_name, "| Rows:", ws.max_row)

        for row_number, row in enumerate(
            ws.iter_rows(min_row=1, max_row=min(ws.max_row, 5),
                         values_only=True), start=1
        ):
            values = [
                str(value).strip() if value is not None else ""
                for value in row
            ]
            if any(values):
                print("ROW", row_number, ":", " | ".join(values[:32]))

    wb.close()

print("\nInspection complete. No workbook was changed.")
